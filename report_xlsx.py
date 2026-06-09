"""
Gera um relatorio XLSX formatado a partir do comparison_report.json.

O JSON continua sendo a fonte de dados completa. Este script cria uma planilha
mais adequada para analise, artigo e revisao manual das divergencias entre
Gemini e Amazon Nova Lite.

Uso:
    python report_xlsx.py
    python report_xlsx.py --input comparison_report.json --output comparison_report.xlsx
"""
import argparse
import json
import math
import os
import statistics
from datetime import datetime
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - mensagem operacional
    OPENPYXL_IMPORT_ERROR = exc
else:
    OPENPYXL_IMPORT_ERROR = None


PROVIDERS = ("gemini", "nova")
PRESENCE_FIELDS = ("cisto", "nodulo", "calcificacao", "microcalcificacao")
COMPARE_FIELDS = PRESENCE_FIELDS + ("bi_rads",)
FIELD_LABELS = {
    "cisto": "Cisto",
    "nodulo": "Nodulo",
    "calcificacao": "Calcificacao",
    "microcalcificacao": "Microcalcificacao",
    "bi_rads": "BI-RADS",
}
PROVIDER_LABELS = {"gemini": "Gemini", "nova": "Nova Lite"}


COLORS = {
    "navy": "17365D",
    "blue": "1F4E78",
    "light_blue": "D9EAF7",
    "green": "D9EAD3",
    "green_text": "274E13",
    "red": "F4CCCC",
    "red_text": "990000",
    "yellow": "FFF2CC",
    "gray": "F3F6F8",
    "dark_gray": "666666",
    "white": "FFFFFF",
}

THIN_BORDER = Border(
    left=Side(style="thin", color="D9E2F3"),
    right=Side(style="thin", color="D9E2F3"),
    top=Side(style="thin", color="D9E2F3"),
    bottom=Side(style="thin", color="D9E2F3"),
)


def load_report(path: str) -> dict[str, Any]:
    """Carrega o JSON de comparacao."""
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"Relatorio JSON nao encontrado: {path}. "
            "Execute `python compare.py` antes de gerar o XLSX."
        )
    with open(path, "r", encoding="utf-8") as file:
        return json.load(file)


def generate_xlsx_report(input_path: str, output_path: str) -> str:
    """
    Gera o XLSX e retorna o caminho do arquivo criado.

    Args:
        input_path: Caminho do comparison_report.json.
        output_path: Caminho do XLSX de saida.
    """
    if OPENPYXL_IMPORT_ERROR:
        raise RuntimeError(
            "Dependencia ausente: openpyxl. "
            "Instale com `pip install -r requirements.txt`."
        ) from OPENPYXL_IMPORT_ERROR

    report = load_report(input_path)

    workbook = Workbook()
    workbook.properties.title = "ClickPalm-IA - Comparacao de Modelos"
    workbook.properties.subject = "Gemini vs Amazon Nova Lite"
    workbook.properties.creator = "ClickPalm-IA"

    ws_summary = workbook.active
    ws_summary.title = "Resumo"
    _build_summary_sheet(ws_summary, report)
    _build_timing_sheet(workbook.create_sheet("Tempos"), report)
    _build_runs_sheet(workbook.create_sheet("Execucoes"), report)
    _build_agreement_sheet(workbook.create_sheet("Concordancia"), report)
    _build_differences_sheet(workbook.create_sheet("Divergencias"), report)
    _build_extractions_sheet(workbook.create_sheet("Extracoes"), report)
    _build_methodology_sheet(workbook.create_sheet("Metodologia"), input_path)

    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False

    workbook.save(output_path)
    return output_path


def _build_summary_sheet(ws, report: dict[str, Any]) -> None:
    ws.sheet_properties.tabColor = COLORS["navy"]
    ws.freeze_panes = "A9"

    _title(ws, "ClickPalm-IA - Relatorio de Comparacao", "Gemini vs Amazon Nova Lite")

    metadata = [
        ("Gerado em", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Diretorio de exames", report.get("exams_dir", "")),
        ("Execucoes por modelo/laudo", report.get("runs_per_model", "")),
        ("Modelo Gemini", report.get("gemini_model", "")),
        ("Modelo Nova", report.get("nova_model", "")),
        ("Regiao AWS", report.get("aws_region", "")),
    ]
    row = 4
    for label, value in metadata:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1
    _style_key_value(ws, 4, row - 1)

    timing_start = row + 1
    ws.cell(timing_start, 1, "Tempo de resposta (ms)")
    _section_header(ws, timing_start, 1, 9)
    headers = [
        "Modelo",
        "Chamadas",
        "Media",
        "Mediana",
        "Min",
        "Max",
        "Total (s)",
        "Media/laudo",
        "Mediana/laudo",
    ]
    _write_row(ws, timing_start + 1, headers, header=True)

    timing = report.get("timing_ms", {})
    for idx, provider in enumerate(PROVIDERS, start=timing_start + 2):
        stats = timing.get(provider, {})
        ws.cell(idx, 1, PROVIDER_LABELS[provider])
        ws.cell(idx, 2, stats.get("count"))
        ws.cell(idx, 3, stats.get("mean"))
        ws.cell(idx, 4, stats.get("median"))
        ws.cell(idx, 5, stats.get("min"))
        ws.cell(idx, 6, stats.get("max"))
        ws.cell(idx, 7, _safe_number(stats.get("total")) / 1000 if stats else None)
        ws.cell(idx, 8, stats.get("mean_per_exam"))
        ws.cell(idx, 9, stats.get("median_per_exam"))
        _style_data_row(ws, idx, 1, 9)

    for data_row in range(timing_start + 2, timing_start + 4):
        for col in range(3, 10):
            ws.cell(data_row, col).number_format = "0.0"
    ws.cell(timing_start + 2, 7).number_format = "0.00"
    ws.cell(timing_start + 3, 7).number_format = "0.00"

    fastest_row = timing_start + 5
    fastest, ratio = _fastest_by_median(timing)
    ws.cell(fastest_row, 1, "Mais rapido por mediana")
    ws.cell(fastest_row, 2, fastest or "n/a")
    ws.cell(fastest_row + 1, 1, "Razao entre medianas")
    ws.cell(fastest_row + 1, 2, ratio)
    ws.cell(fastest_row + 1, 2).number_format = '0.00"x"'
    _style_key_value(ws, fastest_row, fastest_row + 1)

    agreement_start = fastest_row + 3
    agreement = report.get("agreement", {})
    total_matches = _safe_number(agreement.get("total_matches"))
    total_fields = _safe_number(agreement.get("total_fields"))
    pct = total_matches / total_fields if total_fields else None

    ws.cell(agreement_start, 1, "Concordancia geral")
    _section_header(ws, agreement_start, 1, 4)
    general_rows = [
        ("Laudos comparados", agreement.get("exams_compared", 0)),
        ("Campos concordantes", total_matches),
        ("Total de campos", total_fields),
        ("Percentual geral", pct),
    ]
    for offset, (label, value) in enumerate(general_rows, start=1):
        ws.cell(agreement_start + offset, 1, label)
        ws.cell(agreement_start + offset, 2, value)
    ws.cell(agreement_start + 4, 2).number_format = "0.0%"
    _style_key_value(ws, agreement_start + 1, agreement_start + len(general_rows))

    field_start = agreement_start + len(general_rows) + 2
    ws.cell(field_start, 1, "Concordancia por campo")
    _section_header(ws, field_start, 1, 4)
    _write_row(ws, field_start + 1, ["Campo", "Iguais", "Total", "%"], header=True)
    by_field = agreement.get("by_field", {})
    for idx, field in enumerate(COMPARE_FIELDS, start=field_start + 2):
        field_stats = by_field.get(field, {})
        matches = _safe_number(field_stats.get("matches"))
        total = _safe_number(field_stats.get("of"))
        ws.cell(idx, 1, FIELD_LABELS[field])
        ws.cell(idx, 2, matches)
        ws.cell(idx, 3, total)
        ws.cell(idx, 4, matches / total if total else None)
        ws.cell(idx, 4).number_format = "0.0%"
        _style_data_row(ws, idx, 1, 4)

    _add_summary_charts(ws, timing_start, field_start)
    _autosize(ws, max_width=44)


def _build_timing_sheet(ws, report: dict[str, Any]) -> None:
    ws.sheet_properties.tabColor = COLORS["blue"]
    ws.freeze_panes = "A2"

    headers = [
        "Laudo",
        "Gemini media ms",
        "Gemini mediana ms",
        "Gemini min ms",
        "Gemini max ms",
        "Gemini chamadas",
        "Nova media ms",
        "Nova mediana ms",
        "Nova min ms",
        "Nova max ms",
        "Nova chamadas",
        "Diferenca media Nova-Gemini ms",
        "Mais rapido",
        "Razao media",
    ]
    _write_row(ws, 1, headers, header=True)

    for row, entry in enumerate(report.get("results", []), start=2):
        g_times = _timings(entry, "gemini")
        n_times = _timings(entry, "nova")
        g_mean = _mean_or_none(g_times)
        n_mean = _mean_or_none(n_times)
        fastest, ratio = _fastest_from_means(g_mean, n_mean)

        values = [
            entry.get("file", ""),
            g_mean,
            _median_or_none(g_times),
            min(g_times) if g_times else None,
            max(g_times) if g_times else None,
            len(g_times) if g_times else 0,
            n_mean,
            _median_or_none(n_times),
            min(n_times) if n_times else None,
            max(n_times) if n_times else None,
            len(n_times) if n_times else 0,
            n_mean - g_mean if g_mean is not None and n_mean is not None else None,
            fastest,
            ratio,
        ]
        _write_row(ws, row, values)
        _style_data_row(ws, row, 1, len(headers))
        if fastest == "Gemini":
            _fill_range(ws, row, 13, row, 13, COLORS["light_blue"])
        elif fastest == "Nova Lite":
            _fill_range(ws, row, 13, row, 13, COLORS["green"])

    for row in range(2, ws.max_row + 1):
        for col in (2, 3, 4, 5, 7, 8, 9, 10, 12):
            ws.cell(row, col).number_format = "0.0"
        ws.cell(row, 14).number_format = '0.00"x"'

    _add_table(ws, "TabelaTempos")
    _autosize(ws, max_width=36)


def _build_runs_sheet(ws, report: dict[str, Any]) -> None:
    ws.sheet_properties.tabColor = COLORS["blue"]
    ws.freeze_panes = "A2"

    headers = [
        "Laudo",
        "Execucao",
        "Gemini ms",
        "Nova ms",
        "Diferenca Nova-Gemini ms",
        "Mais rapido",
    ]
    _write_row(ws, 1, headers, header=True)

    row = 2
    for entry in report.get("results", []):
        g_times = _timings(entry, "gemini")
        n_times = _timings(entry, "nova")
        max_runs = max(len(g_times), len(n_times), 1)

        for idx in range(max_runs):
            g_time = g_times[idx] if idx < len(g_times) else None
            n_time = n_times[idx] if idx < len(n_times) else None
            fastest, _ratio = _fastest_from_means(g_time, n_time)
            values = [
                entry.get("file", ""),
                idx + 1,
                g_time,
                n_time,
                n_time - g_time if g_time is not None and n_time is not None else None,
                fastest,
            ]
            _write_row(ws, row, values)
            _style_data_row(ws, row, 1, len(headers))
            row += 1

    for data_row in range(2, ws.max_row + 1):
        for col in (3, 4, 5):
            ws.cell(data_row, col).number_format = "0.0"

    _add_table(ws, "TabelaExecucoes")
    _autosize(ws, max_width=32)


def _build_agreement_sheet(ws, report: dict[str, Any]) -> None:
    ws.sheet_properties.tabColor = COLORS["blue"]
    ws.freeze_panes = "A2"

    headers = ["Laudo", "Campo", "Gemini", "Nova Lite", "Concorda"]
    _write_row(ws, 1, headers, header=True)

    row = 2
    for entry in report.get("results", []):
        agreement = entry.get("agreement", {})
        fields = agreement.get("fields", {})
        for field in COMPARE_FIELDS:
            info = fields.get(field, {})
            match = info.get("match")
            values = [
                entry.get("file", ""),
                FIELD_LABELS[field],
                _display_value(info.get("a")),
                _display_value(info.get("b")),
                "Sim" if match is True else "Nao" if match is False else "n/a",
            ]
            _write_row(ws, row, values)
            _style_data_row(ws, row, 1, len(headers))
            if match is True:
                _fill_range(ws, row, 5, row, 5, COLORS["green"])
                ws.cell(row, 5).font = Font(color=COLORS["green_text"], bold=True)
            elif match is False:
                _fill_range(ws, row, 5, row, 5, COLORS["red"])
                ws.cell(row, 5).font = Font(color=COLORS["red_text"], bold=True)
            row += 1

    _add_table(ws, "TabelaConcordancia")
    _autosize(ws, max_width=32)


def _build_differences_sheet(ws, report: dict[str, Any]) -> None:
    ws.sheet_properties.tabColor = COLORS["red"]
    ws.freeze_panes = "A2"

    headers = [
        "Laudo",
        "Campo",
        "Gemini valor",
        "Nova valor",
        "Gemini detalhes",
        "Nova detalhes",
        "Observacao",
    ]
    _write_row(ws, 1, headers, header=True)

    row = 2
    for entry in report.get("results", []):
        for provider in PROVIDERS:
            error = entry.get(provider, {}).get("error")
            if error:
                values = [
                    entry.get("file", ""),
                    "Erro",
                    error if provider == "gemini" else "",
                    error if provider == "nova" else "",
                    "",
                    "",
                    f"Falha na extracao do provedor {PROVIDER_LABELS[provider]}",
                ]
                _write_row(ws, row, values)
                _style_data_row(ws, row, 1, len(headers))
                _fill_range(ws, row, 1, row, len(headers), COLORS["yellow"])
                row += 1

        agreement = entry.get("agreement", {})
        fields = agreement.get("fields", {})
        g_analysis = _analysis(entry, "gemini")
        n_analysis = _analysis(entry, "nova")
        for field in COMPARE_FIELDS:
            info = fields.get(field, {})
            if info.get("match") is not False:
                continue
            values = [
                entry.get("file", ""),
                FIELD_LABELS[field],
                _display_value(info.get("a")),
                _display_value(info.get("b")),
                _field_details(g_analysis, field),
                _field_details(n_analysis, field),
                "Divergencia de extracao entre modelos",
            ]
            _write_row(ws, row, values)
            _style_data_row(ws, row, 1, len(headers))
            _fill_range(ws, row, 1, row, len(headers), COLORS["red"])
            row += 1

    if row == 2:
        _write_row(ws, row, ["Sem divergencias nos campos comparados.", "", "", "", "", "", ""])
        _style_data_row(ws, row, 1, len(headers))
        _fill_range(ws, row, 1, row, len(headers), COLORS["green"])

    _add_table(ws, "TabelaDivergencias")
    _autosize(ws, max_width=56)


def _build_extractions_sheet(ws, report: dict[str, Any]) -> None:
    ws.sheet_properties.tabColor = COLORS["blue"]
    ws.freeze_panes = "A2"

    headers = [
        "Laudo",
        "Provedor",
        "Cisto presente",
        "Cisto detalhes",
        "Nodulo presente",
        "Nodulo detalhes",
        "Calcificacao presente",
        "Calcificacao detalhes",
        "Microcalcificacao presente",
        "Microcalcificacao detalhes",
        "BI-RADS",
        "Outras citacoes",
        "Erro",
    ]
    _write_row(ws, 1, headers, header=True)

    row = 2
    for entry in report.get("results", []):
        for provider in PROVIDERS:
            provider_data = entry.get(provider, {})
            analysis = provider_data.get("analysis") or {}
            values = [
                entry.get("file", ""),
                PROVIDER_LABELS[provider],
                _presence_text(analysis, "cisto"),
                _field_details(analysis, "cisto"),
                _presence_text(analysis, "nodulo"),
                _field_details(analysis, "nodulo"),
                _presence_text(analysis, "calcificacao"),
                _field_details(analysis, "calcificacao"),
                _presence_text(analysis, "microcalcificacao"),
                _field_details(analysis, "microcalcificacao"),
                analysis.get("bi_rads", ""),
                analysis.get("outras_citacoes", ""),
                provider_data.get("error", ""),
            ]
            _write_row(ws, row, values)
            _style_data_row(ws, row, 1, len(headers))
            row += 1

    _add_table(ws, "TabelaExtracoes")
    _autosize(ws, max_width=56)


def _build_methodology_sheet(ws, input_path: str) -> None:
    ws.sheet_properties.tabColor = COLORS["dark_gray"]
    _title(ws, "Metodologia do Relatorio", "Notas para interpretacao dos resultados")

    notes = [
        ("Fonte de dados", input_path),
        (
            "Comparacao",
            "O relatorio compara as saidas dos modelos entre si. Nao ha gabarito "
            "clinico externo neste arquivo.",
        ),
        (
            "Campos avaliados",
            "Cisto, nodulo, calcificacao, microcalcificacao e BI-RADS.",
        ),
        (
            "Concordancia",
            "Um campo concorda quando Gemini e Nova Lite retornam o mesmo valor "
            "normalizado para aquele item.",
        ),
        (
            "Tempo",
            "Os tempos em milissegundos medem somente a chamada remota ao modelo, "
            "conforme registrado pelo compare.py.",
        ),
        (
            "BI-RADS",
            "Na comparacao, o BI-RADS e normalizado para digitos quando possivel "
            "(por exemplo, 'BI-RADS 4C' e tratado como '4').",
        ),
        (
            "Uso em artigo",
            "Use a aba Resumo para metricas agregadas, Tempos/Execucoes para "
            "desempenho e Divergencias/Extracoes para analise qualitativa.",
        ),
    ]

    start_row = 4
    for idx, (label, text) in enumerate(notes, start=start_row):
        ws.cell(idx, 1, label)
        ws.cell(idx, 2, text)
        ws.cell(idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
    _style_key_value(ws, start_row, start_row + len(notes) - 1)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100


def _title(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(size=18, bold=True, color=COLORS["white"])
    ws["A1"].fill = PatternFill("solid", fgColor=COLORS["navy"])
    ws["A1"].alignment = Alignment(horizontal="left")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)

    ws["A2"] = subtitle
    ws["A2"].font = Font(size=11, italic=True, color=COLORS["dark_gray"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)
    ws.row_dimensions[1].height = 28


def _section_header(ws, row: int, start_col: int, end_col: int) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=COLORS["blue"])
        cell.font = Font(color=COLORS["white"], bold=True)
        cell.border = THIN_BORDER
    ws.cell(row, start_col).alignment = Alignment(horizontal="left")


def _write_row(ws, row: int, values: list[Any], header: bool = False) -> None:
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, value)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if header:
            cell.fill = PatternFill("solid", fgColor=COLORS["blue"])
            cell.font = Font(color=COLORS["white"], bold=True)


def _style_key_value(ws, start_row: int, end_row: int) -> None:
    for row in range(start_row, end_row + 1):
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=COLORS["gray"])
        ws.cell(row, 1).font = Font(bold=True, color=COLORS["navy"])
        for col in (1, 2):
            ws.cell(row, col).border = THIN_BORDER
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)


def _style_data_row(ws, row: int, start_col: int, end_col: int) -> None:
    fill = PatternFill("solid", fgColor=COLORS["gray"]) if row % 2 == 0 else None
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if fill:
            cell.fill = fill


def _fill_range(ws, start_row: int, start_col: int, end_row: int, end_col: int, color: str) -> None:
    fill = PatternFill("solid", fgColor=color)
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            ws.cell(row, col).fill = fill


def _add_table(ws, name: str) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _add_summary_charts(ws, timing_start: int, field_start: int) -> None:
    timing_chart = BarChart()
    timing_chart.title = "Mediana de tempo por modelo"
    timing_chart.y_axis.title = "ms"
    timing_chart.x_axis.title = "Modelo"
    data = Reference(ws, min_col=4, min_row=timing_start + 1, max_row=timing_start + 3)
    cats = Reference(ws, min_col=1, min_row=timing_start + 2, max_row=timing_start + 3)
    timing_chart.add_data(data, titles_from_data=True)
    timing_chart.set_categories(cats)
    timing_chart.height = 7
    timing_chart.width = 14
    ws.add_chart(timing_chart, "K4")

    agreement_chart = BarChart()
    agreement_chart.title = "Concordancia por campo"
    agreement_chart.y_axis.title = "%"
    agreement_chart.x_axis.title = "Campo"
    data = Reference(ws, min_col=4, min_row=field_start + 1, max_row=field_start + 6)
    cats = Reference(ws, min_col=1, min_row=field_start + 2, max_row=field_start + 6)
    agreement_chart.add_data(data, titles_from_data=True)
    agreement_chart.set_categories(cats)
    agreement_chart.height = 7
    agreement_chart.width = 14
    ws.add_chart(agreement_chart, "K19")

    data_range = f"D{field_start + 2}:D{field_start + 6}"
    ws.conditional_formatting.add(
        data_range,
        CellIsRule(operator="lessThan", formula=["1"], fill=PatternFill("solid", fgColor=COLORS["yellow"])),
    )


def _autosize(ws, max_width: int = 48) -> None:
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        width = 10
        for cell in column_cells:
            if cell.value is None:
                continue
            value = str(cell.value)
            longest_line = max(len(part) for part in value.splitlines()) if value else 0
            width = max(width, min(longest_line + 2, max_width))
        ws.column_dimensions[column_letter].width = width


def _timings(entry: dict[str, Any], provider: str) -> list[float]:
    provider_data = entry.get(provider, {})
    values = provider_data.get("timings_ms")
    if isinstance(values, list):
        return [_safe_number(value) for value in values if _is_number(value)]
    mean_ms = provider_data.get("mean_ms")
    return [_safe_number(mean_ms)] if _is_number(mean_ms) else []


def _analysis(entry: dict[str, Any], provider: str) -> dict[str, Any]:
    return entry.get(provider, {}).get("analysis") or {}


def _presence_text(analysis: dict[str, Any], field: str) -> str:
    value = analysis.get(field)
    if isinstance(value, dict):
        present = value.get("presente")
        if present is True:
            return "Presente"
        if present is False:
            return "Ausente"
    return ""


def _field_details(analysis: dict[str, Any], field: str) -> str:
    if field == "bi_rads":
        return str(analysis.get("bi_rads", ""))
    value = analysis.get(field)
    if isinstance(value, dict):
        return str(value.get("detalhes", ""))
    return ""


def _display_value(value: Any) -> str:
    if value is True:
        return "Presente"
    if value is False:
        return "Ausente"
    if value is None:
        return "n/a"
    return str(value)


def _mean_or_none(values: list[float]) -> float | None:
    return statistics.mean(values) if values else None


def _median_or_none(values: list[float]) -> float | None:
    return statistics.median(values) if values else None


def _fastest_by_median(timing: dict[str, Any]) -> tuple[str | None, float | None]:
    g = timing.get("gemini", {}).get("median")
    n = timing.get("nova", {}).get("median")
    return _fastest_from_means(_safe_number(g) if _is_number(g) else None, _safe_number(n) if _is_number(n) else None)


def _fastest_from_means(gemini_value: float | None, nova_value: float | None) -> tuple[str | None, float | None]:
    if gemini_value is None or nova_value is None:
        return None, None
    if gemini_value == 0 or nova_value == 0:
        return None, None
    if gemini_value <= nova_value:
        return "Gemini", nova_value / gemini_value
    return "Nova Lite", gemini_value / nova_value


def _safe_number(value: Any) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return 0.0
    return 0.0 if math.isnan(number) or math.isinf(number) else number


def _is_number(value: Any) -> bool:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return False
    return not math.isnan(number) and not math.isinf(number)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Gera XLSX formatado a partir do comparison_report.json."
    )
    parser.add_argument(
        "--input",
        default="comparison_report.json",
        help="Relatorio JSON de entrada (padrao: comparison_report.json).",
    )
    parser.add_argument(
        "--output",
        default="comparison_report.xlsx",
        help="Arquivo XLSX de saida (padrao: comparison_report.xlsx).",
    )
    args = parser.parse_args()

    output = generate_xlsx_report(args.input, args.output)
    print(f"Relatorio XLSX salvo em: {output}")


if __name__ == "__main__":
    main()
